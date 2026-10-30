import csv
import hashlib
import io
import zipfile
from django import forms
from django.contrib import admin, messages
from django.contrib.admin.views.decorators import staff_member_required
from django.core.paginator import Paginator
from django.db import transaction
from django.shortcuts import redirect, render
from django.urls import path, reverse
from openpyxl import load_workbook

from .forms import DatasetUploadForm
from .management.commands.sync_open_data import decimal_value, first_value, make_record
from .management.commands.sync_seoul_archives import normalize_csv_row
from .models import SeoulGridCell, SeoulOpenRecord


DATASET_UPLOADS = {
    "living": {"label": "내국인 생활인구(250m)", "model": SeoulOpenRecord.Dataset.LIVING},
    "area": {"label": "상권영역", "model": SeoulOpenRecord.Dataset.TRADE_AREA},
    "store": {"label": "점포-상권", "model": SeoulOpenRecord.Dataset.STORE},
    "sales": {"label": "추정매출-상권", "model": SeoulOpenRecord.Dataset.SALES},
}


def _to_record(dataset, source):
    row = normalize_csv_row(source)
    if dataset == "area":
        code = str(first_value(row, "TRDAR_CD", "상권_코드"))
        name = str(first_value(row, "TRDAR_CD_NM", "상권_코드_명"))
        key = code or hashlib.sha256(repr(sorted(row.items())).encode("utf-8")).hexdigest()
        record = SeoulOpenRecord(
            dataset=SeoulOpenRecord.Dataset.TRADE_AREA,
            record_key=hashlib.sha256(key.encode("utf-8")).hexdigest(),
            trade_code=code,
            trade_name=name,
            district_name=str(first_value(row, "SIGNGU_CD_NM", "자치구_코드_명")),
            metric_value=decimal_value(first_value(row, "AREA", "AREA_VALUE", "AREA_M2", "영역_면적")),
            metric_label="영역 면적(㎡)",
            x_coordinate=decimal_value(first_value(row, "XCNTS_VALUE", "엑스좌표_값")),
            y_coordinate=decimal_value(first_value(row, "YDNTS_VALUE", "와이좌표_값")),
            payload=row,
        )
        return record

    record = make_record(dataset, row)
    record.payload = row
    if dataset == "living":
        record.period = str(row.get("일자", row.get("YMD", record.period)))
        record.cell_id = str(row.get("250M격자", row.get("CELL_ID", record.cell_id)))
        record.record_key = hashlib.sha256(
            "|".join((record.period, str(row.get("시간", row.get("TT", ""))), str(row.get("행정동코드", row.get("H_DNG_CD", ""))), record.cell_id)).encode("utf-8")
        ).hexdigest()
    elif dataset in ("store", "sales"):
        service_code = str(row.get("SVC_INDUTY_CD", row.get("서비스_업종_코드", "")))
        record.record_key = hashlib.sha256("|".join((record.period, record.trade_code, service_code)).encode("utf-8")).hexdigest()
    return record


def _read_rows(file_handle, filename):
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension == "xlsx":
        workbook = load_workbook(file_handle, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            workbook.close()
            return
        names = [str(value or "").strip() for value in headers]
        try:
            for values in iterator:
                yield dict(zip(names, values))
        finally:
            workbook.close()
        return

    if extension == "zip":
        try:
            archive = zipfile.ZipFile(file_handle)
        except zipfile.BadZipFile as exc:
            raise forms.ValidationError("올바른 ZIP 파일이 아닙니다.") from exc
        entries = [item for item in archive.infolist() if item.filename.lower().endswith((".csv", ".txt"))]
        if not entries:
            archive.close()
            raise forms.ValidationError("ZIP 안에서 CSV/TXT 파일을 찾지 못했습니다.")
        raw = archive.open(entries[0])
        buffered = io.BufferedReader(raw)
        try:
            sample = buffered.peek(8192)
            try:
                sample.decode("utf-8-sig")
                encoding = "utf-8-sig"
            except UnicodeDecodeError:
                encoding = "cp949"
            text = io.TextIOWrapper(buffered, encoding=encoding, errors="replace", newline="")
            yield from csv.DictReader(text)
        finally:
            archive.close()
        return

    uploaded = io.TextIOWrapper(file_handle, encoding="utf-8-sig", errors="replace", newline="")
    try:
        yield from csv.DictReader(uploaded)
    finally:
        uploaded.detach()


@staff_member_required
def dataset_upload_view(request, dataset):
    if dataset not in DATASET_UPLOADS:
        return redirect("admin:index")
    config = DATASET_UPLOADS[dataset]
    form = DatasetUploadForm(request.POST or None, request.FILES or None)
    imported = 0
    if request.method == "POST" and form.is_valid():
        uploaded = form.cleaned_data["file"]
        batch = []
        try:
            with transaction.atomic():
                for source in _read_rows(uploaded, uploaded.name):
                    record = _to_record(dataset, source)
                    batch.append(record)
                    if len(batch) >= 500:
                        imported += _save_upload_batch(batch, form.cleaned_data["replace_existing"])
                        batch = []
                if batch:
                    imported += _save_upload_batch(batch, form.cleaned_data["replace_existing"])
        except (csv.Error, UnicodeError, ValueError, zipfile.BadZipFile, forms.ValidationError) as exc:
            form.add_error("file", str(exc))
        else:
            messages.success(request, f"{config['label']}: {imported:,}행을 처리했습니다.")
            return redirect(reverse("admin:dataset_upload", kwargs={"dataset": dataset}))
    count = SeoulOpenRecord.objects.filter(dataset=config["model"]).count()
    return render(request, "admin/transactions/dataset_upload.html", {
        **admin.site.each_context(request),
        "title": f"{config['label']} 엑셀 업로드",
        "dataset": dataset,
        "label": config["label"],
        "form": form,
        "count": count,
    })


def _save_upload_batch(records, replace_existing):
    if replace_existing:
        with transaction.atomic():
            for record in records:
                SeoulOpenRecord.objects.update_or_create(
                    dataset=record.dataset,
                    record_key=record.record_key,
                    defaults={
                        "period": record.period,
                        "cell_id": record.cell_id,
                        "trade_code": record.trade_code,
                        "trade_name": record.trade_name,
                        "district_name": record.district_name,
                        "category_name": record.category_name,
                        "metric_value": record.metric_value,
                        "metric_label": record.metric_label,
                        "x_coordinate": record.x_coordinate,
                        "y_coordinate": record.y_coordinate,
                        "payload": record.payload,
                    },
                )
        return len(records)
    SeoulOpenRecord.objects.bulk_create(records, batch_size=500, ignore_conflicts=True)
    return len(records)


@admin.register(SeoulOpenRecord)
class SeoulOpenRecordAdmin(admin.ModelAdmin):
    list_display = ("dataset", "period", "trade_name", "cell_id", "category_name", "metric_value", "imported_at")
    list_filter = ("dataset", "period")
    search_fields = ("trade_name", "trade_code", "cell_id", "category_name", "district_name")
    readonly_fields = ("payload", "imported_at")

    def get_urls(self):
        custom_urls = [
            path("upload/<slug:dataset>/", self.admin_site.admin_view(dataset_upload_view), name="dataset_upload"),
        ]
        return custom_urls + super().get_urls()

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context["dataset_uploads"] = [
            {"slug": slug, "label": item["label"]}
            for slug, item in DATASET_UPLOADS.items()
        ]
        return super().changelist_view(request, extra_context=extra_context)


@admin.register(SeoulGridCell)
class SeoulGridCellAdmin(admin.ModelAdmin):
    list_display = ("cell_id", "administrative_code", "imported_at")
    search_fields = ("cell_id", "administrative_code")
    readonly_fields = ("geometry", "properties", "imported_at")